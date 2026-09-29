"""
店舗ごとに形式がバラバラな売上Excelを読み込み、
整形・集計して月次レポート（Excel）を自動作成するスクリプト。

使い方：input/ に各店舗の売上ファイルを置いて実行
    python monthly_report.py
"""
import re
import unicodedata
from pathlib import Path

import pandas as pd
from openpyxl import load_workbook
from openpyxl.chart import BarChart, Reference
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side

INPUT_DIR = Path("input")
OUTPUT = Path("output") / "月次売上レポート_2026Q3.xlsx"

# ---- 設定：ここを書き換えれば別の会社のデータにも対応できる ----
COLUMN_ALIASES = {
    "日付": ["日付", "売上日", "date"],
    "商品名": ["商品名", "品名", "item"],
    "数量": ["数量", "個数", "qty"],
    "単価": ["単価", "unitprice"],
    "売上": ["売上", "金額", "amount"],
}
STORE_NAMES = {"新宿店": "新宿店", "吉祥寺店": "吉祥寺店", "tachikawa": "立川店"}
DEFAULT_YEAR = 2026
FONT = "Yu Gothic"
# ------------------------------------------------------------------

log = []  # 処理ログ（何をどう直したかを記録して、あとから検証できるようにする）


def norm(text):
    """全角・半角のゆれ（全角数字、半角カナなど）と前後の空白を統一する"""
    if text is None or (isinstance(text, float) and pd.isna(text)):
        return ""
    return unicodedata.normalize("NFKC", str(text)).strip()


def to_number(v):
    """「1,350円」「１２」などを数値にする。変換できなければNone"""
    s = norm(v).replace(",", "").replace("円", "")
    try:
        return float(s)
    except ValueError:
        return None


def to_date(v):
    """datetime、「2026-07-01」、「7月1日」などを日付にする"""
    if hasattr(v, "year"):
        return pd.Timestamp(v)
    s = norm(v)
    m = re.fullmatch(r"(\d{1,2})月(\d{1,2})日", s)
    if m:
        return pd.Timestamp(DEFAULT_YEAR, int(m[1]), int(m[2]))
    try:
        return pd.Timestamp(s)
    except ValueError:
        return None


def find_header(raw):
    """先頭10行から、列見出しの行を自動で探す"""
    date_words = {w.lower() for w in COLUMN_ALIASES["日付"]}
    for i in range(min(10, len(raw))):
        if any(norm(c).lower() in date_words for c in raw.iloc[i]):
            return i
    return None


def read_sheet(path, sheet, store):
    raw = pd.read_excel(path, sheet_name=sheet, header=None, dtype=object)
    h = find_header(raw)
    if h is None:
        log.append([store, f"{path.name}/{sheet}", "見出し行が見つからず読み飛ばし", 0])
        return pd.DataFrame()
    headers = [norm(c).lower() for c in raw.iloc[h]]
    col = {}
    for std, aliases in COLUMN_ALIASES.items():
        for a in aliases:
            if a.lower() in headers:
                col[std] = headers.index(a.lower())
                break
    body = raw.iloc[h + 1:]

    rows, skipped = [], {"空行": 0, "合計行など日付でない行": 0}
    for _, r in body.iterrows():
        if all(norm(x) == "" for x in r):
            skipped["空行"] += 1
            continue
        d = to_date(r.iloc[col["日付"]])
        if d is None:
            skipped["合計行など日付でない行"] += 1
            continue
        qty = to_number(r.iloc[col["数量"]])
        price = to_number(r.iloc[col["単価"]]) if "単価" in col else None
        amount = to_number(r.iloc[col["売上"]]) if "売上" in col else None
        if amount is None and qty is not None and price is not None:
            amount = qty * price  # 売上列がないファイルは数量×単価で計算
        if price is None and qty:
            price = amount / qty  # 単価列がないファイルは売上÷数量で計算
        rows.append({
            "日付": d, "月": f"{d.year}-{d.month:02d}", "店舗": store,
            "商品名": norm(r.iloc[col["商品名"]]),
            "数量": qty, "単価": price, "売上": amount,
        })
    log.append([store, f"{path.name}/{sheet}", "読み込み", len(rows)])
    for reason, n in skipped.items():
        if n:
            log.append([store, f"{path.name}/{sheet}", f"除外：{reason}", n])
    return pd.DataFrame(rows)


def load_all():
    frames = []
    for path in sorted(INPUT_DIR.glob("*.xlsx")):
        store = next((v for k, v in STORE_NAMES.items() if k in path.stem), path.stem)
        for sheet in pd.ExcelFile(path).sheet_names:
            frames.append(read_sheet(path, sheet, store))
    return pd.concat(frames, ignore_index=True).sort_values(["日付", "店舗"])


# ---------------- Excel出力 ----------------
HEAD_FILL = PatternFill("solid", fgColor="1F4E78")
THIN = Side(style="thin", color="BFBFBF")


def style_table(ws, n_rows, n_cols, start_row=1):
    for c in range(1, n_cols + 1):
        cell = ws.cell(start_row, c)
        cell.font = Font(name=FONT, bold=True, color="FFFFFF")
        cell.fill = HEAD_FILL
        cell.alignment = Alignment(horizontal="center")
    for r in range(start_row, start_row + n_rows + 1):
        for c in range(1, n_cols + 1):
            cell = ws.cell(r, c)
            cell.border = Border(top=THIN, bottom=THIN, left=THIN, right=THIN)
            if r > start_row:
                cell.font = Font(name=FONT)


def build_report(df):
    OUTPUT.parent.mkdir(exist_ok=True)
    with pd.ExcelWriter(OUTPUT, engine="openpyxl") as w:
        pd.DataFrame().to_excel(w, sheet_name="サマリー")
        df.to_excel(w, sheet_name="明細", index=False)
        pd.DataFrame(log, columns=["店舗", "ファイル/シート", "処理", "件数"]).to_excel(
            w, sheet_name="処理ログ", index=False)

    wb = load_workbook(OUTPUT)
    stores = sorted(df["店舗"].unique())
    months = sorted(df["月"].unique())
    products = sorted(df["商品名"].unique())
    n = len(df) + 1  # 明細の最終行

    # 明細シートの体裁
    d = wb["明細"]
    style_table(d, len(df), 7)
    for r in range(2, n + 1):
        d.cell(r, 1).number_format = "yyyy/mm/dd"
        for c in (6, 7):
            d.cell(r, c).number_format = "#,##0"
    for col, wdt in zip("ABCDEFG", (12, 10, 10, 20, 8, 8, 10)):
        d.column_dimensions[col].width = wdt

    # サマリー：明細を参照する数式で集計（データを直せば自動で再計算される）
    s = wb["サマリー"]
    s["A1"] = "月次売上レポート（2026年7月〜9月）"
    s["A1"].font = Font(name=FONT, bold=True, size=14)
    s["A2"] = "※架空データによる見本。金額は円。集計値は「明細」シートを参照する数式です。"
    s["A2"].font = Font(name=FONT, size=9, color="808080")

    # 表1：店舗×月
    top = 4
    s.cell(top, 1, "店舗 / 月")
    for j, m in enumerate(months, 2):
        s.cell(top, j, m)
    s.cell(top, len(months) + 2, "合計")
    for i, st in enumerate(stores, top + 1):
        s.cell(i, 1, st)
        for j, m in enumerate(months, 2):
            s.cell(i, j, f'=SUMIFS(明細!$G$2:$G${n},明細!$C$2:$C${n},$A{i},明細!$B$2:$B${n},{s.cell(top, j).coordinate})')
    tot = top + len(stores) + 1
    s.cell(tot, 1, "全店合計")
    last_c = len(months) + 2
    for j in range(2, last_c + 1):
        L = s.cell(1, j).column_letter
        for i in range(top + 1, tot + 1):
            if j == last_c:
                s.cell(i, j, f"=SUM(B{i}:{s.cell(i, j - 1).column_letter}{i})")
        s.cell(tot, j, f"=SUM({L}{top + 1}:{L}{tot - 1})")
    style_table(s, len(stores) + 1, last_c, top)
    for r in range(top + 1, tot + 1):
        for c in range(2, last_c + 1):
            s.cell(r, c).number_format = "#,##0"
        if r == tot:
            for c in range(1, last_c + 1):
                s.cell(r, c).font = Font(name=FONT, bold=True)

    # 表2：商品×月
    top2 = tot + 3
    s.cell(top2 - 1, 1, "商品別売上").font = Font(name=FONT, bold=True)
    s.cell(top2, 1, "商品 / 月")
    for j, m in enumerate(months, 2):
        s.cell(top2, j, m)
    s.cell(top2, last_c, "合計")
    for i, p in enumerate(products, top2 + 1):
        s.cell(i, 1, p)
        for j in range(2, last_c):
            s.cell(i, j, f'=SUMIFS(明細!$G$2:$G${n},明細!$D$2:$D${n},$A{i},明細!$B$2:$B${n},{s.cell(top2, j).coordinate})')
        s.cell(i, last_c, f"=SUM(B{i}:{s.cell(i, last_c - 1).column_letter}{i})")
    style_table(s, len(products), last_c, top2)
    for r in range(top2 + 1, top2 + len(products) + 1):
        for c in range(2, last_c + 1):
            s.cell(r, c).number_format = "#,##0"

    # 検算：明細の総額と、サマリーの全店合計が一致するか
    chk = top2 + len(products) + 2
    s.cell(chk, 1, "検算（明細の総額）").font = Font(name=FONT, size=9)
    s.cell(chk, 2, f"=SUM(明細!G2:G{n})").number_format = "#,##0"
    s.cell(chk, 3, f'=IF(B{chk}={s.cell(tot, last_c).coordinate},"一致","不一致")')

    s.column_dimensions["A"].width = 20
    for j in range(2, last_c + 1):
        s.column_dimensions[s.cell(1, j).column_letter].width = 14

    # グラフ：店舗別の月次売上
    ch = BarChart()
    ch.title = "店舗別 月次売上"
    ch.y_axis.title = "売上（円）"
    ch.height, ch.width = 8, 16
    data = Reference(s, min_col=2, max_col=len(months) + 1, min_row=top, max_row=tot - 1)
    ch.add_data(data, titles_from_data=True)
    ch.set_categories(Reference(s, min_col=1, min_row=top + 1, max_row=tot - 1))
    s.add_chart(ch, f"{s.cell(1, last_c + 2).column_letter}{top}")

    # 処理ログの体裁
    lg = wb["処理ログ"]
    style_table(lg, len(log), 4)
    for col, wdt in zip("ABCD", (10, 34, 30, 8)):
        lg.column_dimensions[col].width = wdt

    wb.save(OUTPUT)


if __name__ == "__main__":
    df = load_all()
    build_report(df)
    print(f"{len(df)}件の売上明細を整形し、{OUTPUT} を作成しました")
